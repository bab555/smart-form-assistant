import pandas as pd
from docx import Document
import io
import logging

logger = logging.getLogger(__name__)

def parse_file_content(file_bytes: bytes, filename: str) -> str:
    """
    统一文件解析入口
    支持: .xlsx, .xls, .csv, .docx, .txt
    说明：
    - .doc（老 Word）python-docx 无法直接解析，建议用户另存为 .docx 后上传
    """
    filename = filename.lower()
    
    try:
        if filename.endswith(('.xlsx', '.xls', '.csv')):
            return _parse_excel(file_bytes, filename)
        elif filename.endswith('.docx'):
            return _parse_word(file_bytes)
        elif filename.endswith('.doc'):
            return "暂不支持 .doc（老版 Word 格式），请另存为 .docx 后再上传。"
        elif filename.endswith('.txt'):
            return file_bytes.decode('utf-8')
        else:
            return f"不支持的文件格式: {filename}。请上传 Excel, Word, CSV 或 TXT。"
    except Exception as e:
        logger.error(f"文件解析失败: {e}")
        return f"文件解析发生错误: {str(e)}"

def _df_to_simple_text(df):
    """
    将 DataFrame 转换为简单的文本表格（类似 Word 表格提取风格）
    不使用 markdown 对齐，去除多余空格，最接近自然语言
    """
    lines = []
    # 表头
    headers = [str(c).strip() for c in df.columns]
    lines.append("| " + " | ".join(headers) + " |")
    
    # 数据行
    for _, row in df.iterrows():
        # 逐个单元格转字符串，去掉换行符
        cells = [str(val).strip().replace('\n', ' ') for val in row]
        lines.append("| " + " | ".join(cells) + " |")
        
    return "\n".join(lines)


def _parse_excel(file_bytes: bytes, filename: str) -> str:
    """解析 Excel/CSV 为简单文本表格"""
    try:
        if filename.endswith('.csv'):
            df = pd.read_csv(io.BytesIO(file_bytes))
            df = df.fillna('')
            return _df_to_simple_text(df)
        else:
            # 兼容：有些文件扩展名是 .xls，但实际内容是 .xlsx（zip: PK..）
            # xls (OLE2/CFB) header: D0 CF 11 E0 A1 B1 1A E1
            ole2_header = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
            is_zip = file_bytes[:2] == b"PK"
            is_ole2 = file_bytes[:8] == ole2_header

            # 注意：很多 Excel 会有多个 sheet，而 pd.read_excel 默认只读第一个 sheet，
            # 这会导致“看起来读取成功但内容为空/不对”的坑。
            engine = None
            if is_ole2:
                engine = "xlrd"
            elif is_zip:
                engine = "openpyxl"

            # 读取所有 sheet（若只有一个 sheet，则也兼容）
            xls = pd.ExcelFile(io.BytesIO(file_bytes), engine=engine) if engine else pd.ExcelFile(io.BytesIO(file_bytes))
            sheets = xls.sheet_names or []
            if not sheets:
                return "Excel 读取失败: 未发现任何工作表（sheet）"

            parts = []
            for sname in sheets:
                try:
                    df_sheet = xls.parse(sname).fillna('')
                except Exception as e:
                    parts.append(f"### Sheet: {sname}\n读取失败: {str(e)}\n")
                    continue

                # 跳过完全空 sheet
                if df_sheet.empty and len(df_sheet.columns) == 0:
                    continue

                # 控制 token：每个 sheet 最多取前 200 行，避免巨大 Excel 把上下文撑爆
                truncated = False
                if len(df_sheet) > 200:
                    df_sheet = df_sheet.head(200)
                    truncated = True

                # 使用自定义的简单文本转换，替代 to_markdown
                # 这样 Excel 和 Word 的输出格式就高度一致了
                txt = _df_to_simple_text(df_sheet)
                
                if truncated:
                    txt += f"\n\n（Sheet {sname} 已截断，仅展示前 200 行）"

                # 移除 "### Sheet: name" 标题，直接拼接内容，使其与 Word 格式完全一致
                parts.append(f"{txt}\n")

            if not parts:
                return "Excel 读取失败: 所有工作表均为空"

            return "\n".join(parts).strip()
    except Exception as e:
        # 提示：扩展名与真实格式不一致时经常会失败
        msg = str(e)
        hint = ""
        if filename.endswith(".xls") and file_bytes[:2] == b"PK":
            hint = "（提示：该文件扩展名为 .xls，但内容看起来是 .xlsx，请尝试另存为 .xlsx 或改扩展名后再上传）"
        return f"Excel 读取失败: {msg}{hint}"

def _parse_word(file_bytes: bytes) -> str:
    """解析 Word 文档为纯文本"""
    try:
        doc = Document(io.BytesIO(file_bytes))
        full_text = []
        
        # 1. 提取段落
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        
        # 2. 提取表格（非常重要，订单常在表格里）
        if doc.tables:
            full_text.append("\n--- 文档内的表格数据 ---\n")
            for table in doc.tables:
                # 简单转换为管道符分隔的文本 | A | B |
                for row in table.rows:
                    row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    full_text.append("| " + " | ".join(row_cells) + " |")
                full_text.append("") # 表格间空行
                
        return "\n".join(full_text)
    except Exception as e:
        return f"Word 读取失败: {str(e)}"

